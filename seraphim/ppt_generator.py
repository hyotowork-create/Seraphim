"""찬양 가사 PPT와 광고 슬라이드를 생성하는 모듈.

python-pptx 만 사용하여 외부 PowerPoint 설치 없이 .pptx 파일을 만든다.
16:9 와이드 화면, 큰 중앙 정렬 가사, 세련된 그라데이션 배경을 적용한다.
"""

from __future__ import annotations

import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from .themes import Theme, get_theme

# 16:9 와이드 슬라이드 크기 (EMU). 13.333" x 7.5"
SLIDE_W = Emu(12192000)
SLIDE_H = Emu(6858000)

# 한 슬라이드에 들어갈 가사 최대 줄 수 (가독성 기준)
MAX_LINES_PER_SLIDE = 4


def _hex(color: str) -> RGBColor:
    return RGBColor.from_string(color)


def _set_gradient_bg(slide, theme: Theme) -> None:
    """슬라이드 배경을 대각선 그라데이션으로 채운다."""
    fill = slide.background.fill
    fill.gradient()
    # 그라데이션 정지점 2개 설정
    stops = fill.gradient_stops
    stops[0].color.rgb = _hex(theme.bg)
    stops[0].position = 0.0
    stops[1].color.rgb = _hex(theme.bg_gradient)
    stops[1].position = 1.0
    try:
        fill.gradient_angle = 135.0
    except Exception:
        pass  # 일부 환경에서 각도 미지원 시 무시


def _add_textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return box, tf


def _style_run(run, *, size, color, bold=False, font="Noto Sans KR"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _hex(color)
    run.font.name = font
    # 한글 폰트(동아시아 글꼴)도 동일하게 지정
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = rpr.makeelement(qn("a:ea"), {})
        rpr.append(ea)
    ea.set("typeface", font)


def _add_accent_bar(slide, theme: Theme, top, width_emu=Emu(1600000)):
    """제목 아래 들어가는 가는 강조 막대."""
    from pptx.enum.shapes import MSO_SHAPE
    left = Emu(int((SLIDE_W - width_emu) / 2))
    bar = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width_emu, Emu(60000)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = _hex(theme.accent)
    bar.line.fill.background()
    bar.shadow.inherit = False
    return bar


def _blank_slide(prs: Presentation):
    # 레이아웃 6 = 완전 빈 슬라이드
    return prs.slides.add_slide(prs.slide_layouts[6])


def split_lyrics(lyrics: str, max_lines: int = MAX_LINES_PER_SLIDE) -> list[str]:
    """가사 텍스트를 슬라이드 단위 블록으로 나눈다.

    - 빈 줄(`\\n\\n`)이 있으면 그것을 절 구분으로 우선 사용한다.
    - 한 절이 max_lines 보다 길면 max_lines 단위로 추가 분할한다.
    - 빈 줄이 전혀 없으면 max_lines 단위로 균등 분할한다.
    """
    text = lyrics.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text:
        return []

    raw_blocks = re.split(r"\n\s*\n", text)
    blocks: list[str] = []
    for block in raw_blocks:
        lines = [ln.strip() for ln in block.split("\n") if ln.strip()]
        if not lines:
            continue
        for i in range(0, len(lines), max_lines):
            blocks.append("\n".join(lines[i:i + max_lines]))
    return blocks


def _add_title_slide(prs, theme, song_title, subtitle=""):
    slide = _blank_slide(prs)
    _set_gradient_bg(slide, theme)

    _, tf = _add_textbox(slide, Emu(800000), Emu(2400000),
                         Emu(SLIDE_W - 1600000), Emu(1800000))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = song_title
    _style_run(run, size=54, color=theme.text, bold=True, font=theme.title_font)

    _add_accent_bar(slide, theme, Emu(4350000))

    if subtitle:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = subtitle
        _style_run(r2, size=24, color=theme.subtext, font=theme.body_font)
    return slide


def _add_lyric_slide(prs, theme, block, footer=""):
    slide = _blank_slide(prs)
    _set_gradient_bg(slide, theme)

    _, tf = _add_textbox(slide, Emu(700000), Emu(700000),
                         Emu(SLIDE_W - 1400000), Emu(SLIDE_H - 1400000))
    for idx, line in enumerate(block.split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(12)
        run = p.add_run()
        run.text = line
        _style_run(run, size=40, color=theme.text, bold=True, font=theme.body_font)

    if footer:
        # 오른쪽 하단 곡명 표기
        _, ff = _add_textbox(slide, Emu(SLIDE_W - 4000000),
                             Emu(SLIDE_H - 700000), Emu(3700000), Emu(500000))
        ff.vertical_anchor = MSO_ANCHOR.BOTTOM
        fp = ff.paragraphs[0]
        fp.alignment = PP_ALIGN.RIGHT
        fr = fp.add_run()
        fr.text = footer
        _style_run(fr, size=14, color=theme.subtext, font=theme.body_font)
    return slide


def _add_ad_slide(prs, theme, title, content):
    slide = _blank_slide(prs)
    _set_gradient_bg(slide, theme)

    # 제목
    _, tf = _add_textbox(slide, Emu(800000), Emu(700000),
                         Emu(SLIDE_W - 1600000), Emu(1200000))
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = title
    _style_run(run, size=44, color=theme.accent, bold=True, font=theme.title_font)

    _add_accent_bar(slide, theme, Emu(2050000), width_emu=Emu(1200000))

    # 내용 (좌측 정렬, 항목별 줄바꿈)
    _, bf = _add_textbox(slide, Emu(1400000), Emu(2500000),
                         Emu(SLIDE_W - 2800000), Emu(SLIDE_H - 3200000))
    bf.vertical_anchor = MSO_ANCHOR.TOP
    lines = [ln.strip() for ln in content.replace("\r\n", "\n").split("\n") if ln.strip()]
    for idx, line in enumerate(lines):
        p = bf.paragraphs[0] if idx == 0 else bf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(14)
        run = p.add_run()
        run.text = f"•  {line}"
        _style_run(run, size=26, color=theme.text, font=theme.body_font)
    return slide


def build_presentation(data: dict, theme_key: str | None = None) -> Presentation:
    """전체 예배 PPT를 생성한다.

    data 구조::

        {
          "service": {"church": "...", "date": "...", "service_name": "..."},
          "songs": [{"title": "...", "lyrics": "..."}, ...],
          "sermon": {"title": "...", "preacher": "...", "passage": "..."},
          "ads": [{"title": "...", "content": "..."}, ...]
        }
    """
    theme = get_theme(theme_key)
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    service = data.get("service", {})

    # 1) 예배 표지
    _add_title_slide(
        prs, theme,
        service.get("service_name") or "주일예배",
        " · ".join(x for x in [service.get("church"), service.get("date")] if x),
    )

    # 2) 찬양: 곡마다 표지 + 가사 슬라이드
    for song in data.get("songs", []):
        title = (song.get("title") or "").strip()
        lyrics = song.get("lyrics") or ""
        if not title and not lyrics.strip():
            continue
        _add_title_slide(prs, theme, title or "찬양", "찬양")
        for block in split_lyrics(lyrics):
            _add_lyric_slide(prs, theme, block, footer=title)

    # 3) 설교 표지
    sermon = data.get("sermon", {})
    if sermon.get("title"):
        sub = " · ".join(x for x in [sermon.get("passage"), sermon.get("preacher")] if x)
        _add_title_slide(prs, theme, sermon["title"], sub or "설교")

    # 4) 광고 슬라이드
    ads = [a for a in data.get("ads", []) if (a.get("title") or a.get("content"))]
    if ads:
        _add_title_slide(prs, theme, "광고 및 안내", "")
        for ad in ads:
            _add_ad_slide(prs, theme,
                          ad.get("title") or "안내", ad.get("content") or "")

    return prs
